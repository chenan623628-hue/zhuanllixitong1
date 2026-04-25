"""
专利-标准比对系统 V1.0
M04 文件上传与安全网关模块 - Excel 预检服务
"""
import json
import logging
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any

from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.exceptions import FileException
from app.core.schemas import ErrorCode

logger = logging.getLogger(__name__)


MAX_EXCEL_ROWS = 50
MAX_EXCEL_COLUMNS = 20


@dataclass
class ExcelValidationError:
    row: int
    column: int
    column_name: str
    error_type: str
    error_message: str


@dataclass
class ExcelPreviewResult:
    valid: bool
    total_rows: int
    valid_rows: int
    invalid_rows: int
    headers: list[str] = field(default_factory=list)
    sample_data: list[dict[str, Any]] = field(default_factory=list)
    errors: list[ExcelValidationError] = field(default_factory=list)
    warnings: list[dict[str, Any]] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)


class ExcelService:
    def __init__(self, db: Session):
        self.db = db
    
    def preview_excel(
        self,
        file_path: str,
        max_rows: int = MAX_EXCEL_ROWS,
        require_headers: bool = True,
    ) -> ExcelPreviewResult:
        try:
            import openpyxl
            from openpyxl.utils import get_column_letter
        except ImportError:
            logger.warning("openpyxl not installed, using mock Excel preview")
            return self._mock_preview_excel(file_path, max_rows, require_headers)
        
        try:
            wb = openpyxl.load_workbook(
                file_path,
                read_only=True,
                data_only=True,
            )
            
            ws = wb.active
            
            total_rows = 0
            headers = []
            sample_data = []
            errors = []
            warnings = []
            
            current_row = 0
            headers_read = False
            
            for row in ws.iter_rows(values_only=True):
                current_row += 1
                
                if all(cell is None or cell == "" for cell in row):
                    continue
                
                total_rows += 1
                
                if not headers_read and require_headers:
                    headers = self._parse_headers(row)
                    headers_read = True
                    continue
                
                if len(sample_data) < 10:
                    row_data = self._parse_row(row, headers, current_row)
                    sample_data.append(row_data)
                
                row_errors = self._validate_row(row, headers, current_row, total_rows)
                errors.extend(row_errors)
                
                if total_rows > max_rows:
                    warnings.append({
                        "type": "row_limit_exceeded",
                        "message": f"行数超过限制（最大 {max_rows} 行），仅处理前 {max_rows} 行",
                        "actual_rows": total_rows,
                        "max_rows": max_rows,
                    })
                    break
            
            wb.close()
            
            valid_rows = total_rows - len(errors)
            
            result = ExcelPreviewResult(
                valid=len(errors) == 0,
                total_rows=total_rows,
                valid_rows=valid_rows,
                invalid_rows=len(errors),
                headers=headers,
                sample_data=sample_data[:10],
                errors=errors,
                warnings=warnings,
                metadata={
                    "sheet_name": ws.title,
                    "max_columns": len(headers),
                    "max_rows_allowed": max_rows,
                    "parsed_at": datetime.now().isoformat(),
                },
            )
            
            logger.info(f"Excel preview complete: {total_rows} rows, {len(errors)} errors")
            
            return result
            
        except Exception as e:
            logger.error(f"Excel preview failed: {str(e)}", exc_info=True)
            raise FileException(
                code=ErrorCode.FILE_PARSE_ERROR,
                message=f"Excel 文件解析失败: {str(e)}"
            )
    
    def _mock_preview_excel(
        self,
        file_path: str,
        max_rows: int,
        require_headers: bool,
    ) -> ExcelPreviewResult:
        logger.info(f"Using mock Excel preview for: {file_path}")
        
        return ExcelPreviewResult(
            valid=True,
            total_rows=0,
            valid_rows=0,
            invalid_rows=0,
            headers=["列1", "列2", "列3"],
            sample_data=[],
            errors=[],
            warnings=[{
                "type": "mock_preview",
                "message": "openpyxl 未安装，使用模拟预览模式",
            }],
            metadata={
                "mock": True,
                "parsed_at": datetime.now().isoformat(),
            },
        )
    
    def _parse_headers(self, row: tuple) -> list[str]:
        headers = []
        for i, cell in enumerate(row):
            if cell is None:
                headers.append(f"列{i+1}")
            else:
                header = str(cell).strip()
                if not header:
                    headers.append(f"列{i+1}")
                else:
                    headers.append(header)
        return headers
    
    def _parse_row(self, row: tuple, headers: list[str], row_num: int) -> dict[str, Any]:
        result = {}
        for i, cell in enumerate(row):
            if i < len(headers):
                key = headers[i]
            else:
                key = f"列{i+1}"
            
            if isinstance(cell, datetime):
                result[key] = cell.isoformat()
            elif cell is None:
                result[key] = ""
            else:
                result[key] = str(cell)
        
        return result
    
    def _validate_row(
        self,
        row: tuple,
        headers: list[str],
        excel_row_num: int,
        logical_row_num: int,
    ) -> list[ExcelValidationError]:
        errors = []
        
        if not headers:
            return errors
        
        required_columns = [0]
        for col_idx in required_columns:
            if col_idx < len(row):
                cell = row[col_idx]
                if cell is None or str(cell).strip() == "":
                    col_name = headers[col_idx] if col_idx < len(headers) else f"列{col_idx+1}"
                    errors.append(ExcelValidationError(
                        row=logical_row_num,
                        column=col_idx + 1,
                        column_name=col_name,
                        error_type="required",
                        error_message=f"必填列 '{col_name}' 不能为空",
                    ))
        
        return errors
    
    def validate_excel_import(
        self,
        file_path: str,
        schema: dict[str, Any] | None = None,
    ) -> ExcelPreviewResult:
        return self.preview_excel(file_path)
    
    def extract_excel_data(
        self,
        file_path: str,
        start_row: int = 1,
        end_row: int | None = None,
    ) -> list[dict[str, Any]]:
        try:
            import openpyxl
        except ImportError:
            logger.warning("openpyxl not installed, returning empty data")
            return []
        
        try:
            wb = openpyxl.load_workbook(
                file_path,
                read_only=True,
                data_only=True,
            )
            
            ws = wb.active
            
            headers = []
            data = []
            current_row = 0
            
            for row in ws.iter_rows(values_only=True):
                current_row += 1
                
                if current_row < start_row:
                    continue
                
                if end_row and current_row > end_row:
                    break
                
                if all(cell is None or cell == "" for cell in row):
                    continue
                
                if not headers:
                    headers = self._parse_headers(row)
                    continue
                
                row_data = self._parse_row(row, headers, current_row)
                data.append(row_data)
            
            wb.close()
            
            logger.info(f"Extracted {len(data)} rows from Excel")
            
            return data
            
        except Exception as e:
            logger.error(f"Excel data extraction failed: {str(e)}", exc_info=True)
            raise FileException(
                code=ErrorCode.FILE_PARSE_ERROR,
                message=f"Excel 数据提取失败: {str(e)}"
            )
    
    def get_excel_info(self, file_path: str) -> dict[str, Any]:
        try:
            import openpyxl
        except ImportError:
            return {
                "worksheets": [],
                "active_sheet": None,
                "total_rows": 0,
                "total_columns": 0,
                "openpyxl_available": False,
            }
        
        try:
            wb = openpyxl.load_workbook(
                file_path,
                read_only=True,
                data_only=True,
            )
            
            ws = wb.active
            
            row_count = 0
            col_count = 0
            
            for row in ws.iter_rows(values_only=True):
                row_count += 1
                col_count = max(col_count, len(row))
                if row_count > 1000:
                    break
            
            info = {
                "worksheets": wb.sheetnames,
                "active_sheet": ws.title,
                "total_rows": row_count,
                "total_columns": col_count,
                "openpyxl_available": True,
                "parsed_at": datetime.now().isoformat(),
            }
            
            wb.close()
            
            return info
            
        except Exception as e:
            logger.error(f"Excel info extraction failed: {str(e)}")
            return {
                "worksheets": [],
                "active_sheet": None,
                "total_rows": 0,
                "total_columns": 0,
                "error": str(e),
                "openpyxl_available": True,
            }
